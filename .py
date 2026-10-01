import os
import re
from pathlib import Path
from datetime import datetime
from typing import TypedDict

from dotenv import load_dotenv
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape
from langchain_groq import ChatGroq
from langchain_core.messages import HumanMessage, SystemMessage
from langgraph.graph import StateGraph, START, END


# Load credentials from the .env file beside this script, independent of the
# directory from which the script is launched.
load_dotenv(Path(__file__).with_name(".env"))
groq_api_key = os.getenv("GROQ_API_KEY")
if not groq_api_key:
    raise RuntimeError(
        "GROQ_API_KEY is missing. Add it to the .env file beside this script "
        "or set it in your environment."
    )

# Initialize Groq LLM
llm = ChatGroq(
    model="openai/gpt-oss-20b",
    temperature=0.3,
    api_key=groq_api_key
)


# Shared state for all agents
class AgentState(TypedDict):
    question: str
    research: str
    technical: str
    report: str


# Coordinator node
def coordinator(state: AgentState):
    print("\n[Coordinator] Starting workflow")

    return {
        "research": "",
        "technical": "",
        "report": ""
    }


# Agent 1: Research Agent
def research_agent(state: AgentState):
    print("\n[Research Agent] Working...")

    prompt = f"""
    Analyze the following question as a Research Agent.

    Identify:
    - Important concepts
    - Requirements
    - Benefits
    - Challenges

    Question: {state['question']}
    """

    response = llm.invoke([
        SystemMessage(content="You are an IT Research Agent."),
        HumanMessage(content=prompt)
    ])

    return {"research": response.content}


# Agent 2: Technical Agent
def technical_agent(state: AgentState):
    print("\n[Technical Agent] Working...")

    prompt = f"""
    You are a Kubernetes Technical Architect.

    Based on the research, propose a technical solution.

    Include:
    - Architecture
    - Kubernetes components
    - Deployment approach
    - Security
    - Monitoring

    User question:
    {state['question']}

    Research findings:
    {state['research']}
    """

    response = llm.invoke([
        SystemMessage(content="You are a Kubernetes expert."),
        HumanMessage(content=prompt)
    ])

    return {"technical": response.content}


# Agent 3: Report Agent
def report_agent(state: AgentState):
    print("\n[Report Agent] Working...")

    prompt = f"""
    Prepare a structured technical report.

    Include:
    1. Executive summary
    2. Research findings
    3. Technical architecture
    4. Implementation steps
    5. Conclusion

    User question:
    {state['question']}

    Research:
    {state['research']}

    Technical solution:
    {state['technical']}

    Do not invent unsupported facts.
    """

    response = llm.invoke([
        SystemMessage(content="You are a Technical Report Agent."),
        HumanMessage(content=prompt)
    ])

    return {"report": response.content}


# Build LangGraph
graph = StateGraph(AgentState)

# Register nodes
graph.add_node("coordinator", coordinator)
graph.add_node("research", research_agent)
graph.add_node("technical", technical_agent)
graph.add_node("report", report_agent)

# Define workflow
graph.add_edge(START, "coordinator")
graph.add_edge("coordinator", "research")
graph.add_edge("research", "technical")
graph.add_edge("technical", "report")
graph.add_edge("report", END)

# Compile graph
app = graph.compile()


def save_report(report: str, output_format: str) -> Path:
    """Save the generated report as a Word document, PDF, or both."""
    output_dir = Path(__file__).with_name("reports")
    output_dir.mkdir(exist_ok=True)
    filename = f"report_{datetime.now():%Y%m%d_%H%M%S}"
    saved_paths = []

    if output_format in ("word", "both"):
        word_path = output_dir / f"{filename}.docx"
        document = Document()
        document.add_heading("Technical Report", level=0)
        for line in report.splitlines():
            line = line.strip()
            if not line:
                continue
            heading = re.match(r"^#{1,3}\s+(.+)$", line)
            if heading:
                document.add_heading(heading.group(1), level=min(line.count("#"), 3))
            elif line.startswith(("- ", "* ")):
                document.add_paragraph(line[2:], style="List Bullet")
            elif re.match(r"^\d+[.)]\s+", line):
                document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line), style="List Number")
            else:
                document.add_paragraph(line)
        document.save(word_path)
        saved_paths.append(word_path)

    if output_format in ("pdf", "both"):
        pdf_path = output_dir / f"{filename}.pdf"
        styles = getSampleStyleSheet()
        story = [Paragraph("Technical Report", styles["Title"]), Spacer(1, 0.2 * inch)]
        for line in report.splitlines():
            line = line.strip()
            if not line:
                story.append(Spacer(1, 0.08 * inch))
                continue
            heading = re.match(r"^#{1,3}\s+(.+)$", line)
            if heading:
                story.append(Paragraph(escape(heading.group(1)), styles["Heading" + str(min(line.count("#"), 3))]))
            else:
                text = re.sub(r"^[-*]\s+", "• ", line)
                text = re.sub(r"^\d+[.)]\s+", "• ", text)
                story.append(Paragraph(escape(text), styles["BodyText"]))
        SimpleDocTemplate(str(pdf_path), pagesize=letter).build(story)
        saved_paths.append(pdf_path)

    return saved_paths[0] if len(saved_paths) == 1 else output_dir


# Execute application
if __name__ == "__main__":
    question = input("Enter your question: ")
    output_format = input("Save format (word/pdf/both) [word]: ").strip().lower() or "word"
    if output_format not in {"word", "pdf", "both"}:
        raise ValueError("Choose 'word', 'pdf', or 'both'.")

    result = app.invoke({
        "question": question,
        "research": "",
        "technical": "",
        "report": ""
    })

    print("\n========== FINAL REPORT ==========")
    print(result["report"])
    saved_to = save_report(result["report"], output_format)
    print(f"\nReport saved to: {saved_to}")
